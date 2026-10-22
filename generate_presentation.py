# -*- coding: utf-8 -*-
"""
Generate Presentation Deck for KNW PRO Realtor Industrial Project
Format: 16:9 Widescreen Presentation (.pptx)
Canva Compatible: Can be imported directly into Canva via Upload -> Presentation
Module: IT3901 Industrial Training | Lanka Nippon BizTech Institute (LNBTI)
Organization: KNW PRO Realtor (Private) Limited, Colombo, Sri Lanka
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Color Palette ─────────────────────────────────────────────────────────────
COLOR_BG_DARK = RGBColor(15, 23, 42)      # #0F172A Deep Slate Navy
COLOR_BG_LIGHT = RGBColor(248, 250, 252)  # #F8FAFC Soft Off-White
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_CARD_BG = RGBColor(255, 255, 255)
COLOR_CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0
COLOR_TEXT_DARK = RGBColor(15, 23, 42)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
COLOR_TEXT_LIGHT = RGBColor(241, 245, 249)
COLOR_GOLD = RGBColor(217, 119, 6)        # #D97706 Warm Amber/Gold
COLOR_EMERALD = RGBColor(5, 150, 105)     # #059669 Emerald Green
COLOR_BLUE = RGBColor(37, 99, 235)        # #2563EB Primary Blue
COLOR_CARD_DARK = RGBColor(30, 41, 59)    # #1E293B

FONT_TITLE = "Arial"
FONT_BODY = "Arial"

def set_slide_background(slide, color):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = color
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, tag_text, title_text, dark_mode=False):
    # Category Tag / Badge
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.name = FONT_TITLE
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_GOLD if dark_mode else COLOR_EMERALD

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_TITLE
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE if dark_mode else COLOR_TEXT_DARK

def add_card(slide, left, top, width, height, title="", border_color=COLOR_CARD_BORDER, bg_color=COLOR_CARD_BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)
    
    if title:
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_TITLE
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_DARK
        p.space_after = Pt(8)
    return shape, tf

def add_bullet_point(tf, bold_prefix, text, font_size=11, dark_text=True, bullet_char="• "):
    p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
    p.space_before = Pt(4)
    p.space_after = Pt(4)
    p.line_spacing = 1.2
    
    run1 = p.add_run()
    run1.text = bullet_char + (bold_prefix + ": " if bold_prefix else "")
    run1.font.name = FONT_BODY
    run1.font.size = Pt(font_size)
    run1.font.bold = True
    run1.font.color.rgb = (COLOR_GOLD if not dark_text else COLOR_EMERALD) if bold_prefix else (COLOR_TEXT_LIGHT if not dark_text else COLOR_TEXT_DARK)
    
    run2 = p.add_run()
    run2.text = text
    run2.font.name = FONT_BODY
    run2.font.size = Pt(font_size)
    run2.font.color.rgb = COLOR_TEXT_LIGHT if not dark_text else COLOR_TEXT_DARK

def set_speaker_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 1: Title Slide (Dark Navy Hero)
    # ═════════════════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, COLOR_BG_DARK)

    # Decorative top bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_GOLD
    bar.line.fill.background()

    # Badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(4.5), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = COLOR_CARD_DARK
    badge.line.color.rgb = COLOR_GOLD
    b_tf = badge.text_frame
    b_tf.margin_left = Inches(0.2)
    b_tf.margin_top = Inches(0.08)
    bp = b_tf.paragraphs[0]
    bp.text = "LNBTI • IT3901 INDUSTRIAL TRAINING PROJECT"
    bp.font.name = FONT_TITLE
    bp.font.size = Pt(11)
    bp.font.bold = True
    bp.font.color.rgb = COLOR_GOLD

    # Main Title
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(1.8))
    t_tf = t_box.text_frame
    t_tf.word_wrap = True
    p1 = t_tf.paragraphs[0]
    p1.text = "KNW PRO REALTOR"
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_after = Pt(8)

    p2 = t_tf.add_paragraph()
    p2.text = "Design & Implementation of a Digital Property Trading, Title Deed Verification & Legal Advisory Platform"
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLOR_GOLD

    # Meta card (Student & Org Info)
    meta_shape, meta_tf = add_card(s1, Inches(1.0), Inches(4.0), Inches(11.333), Inches(2.5), "", COLOR_CARD_BORDER, COLOR_CARD_DARK)
    
    mp1 = meta_tf.paragraphs[0]
    mp1.text = "Candidate Information & Host Enterprise"
    mp1.font.name = FONT_TITLE
    mp1.font.size = Pt(15)
    mp1.font.bold = True
    mp1.font.color.rgb = COLOR_GOLD
    mp1.space_after = Pt(10)

    add_bullet_point(meta_tf, "Student Name", "[Student's Full Name]", 13, dark_text=False)
    add_bullet_point(meta_tf, "Registration Number", "[Reg Number e.g. DSE/2024/001] | Higher Diploma in Software Engineering", 13, dark_text=False)
    add_bullet_point(meta_tf, "Academic Institute", "Lanka Nippon BizTech Institute (LNBTI) • Academic Year 2026", 13, dark_text=False)
    add_bullet_point(meta_tf, "Host Organization", "KNW PRO Realtor (Private) Limited, IT & Digital Systems Division, Colombo", 13, dark_text=False)
    add_bullet_point(meta_tf, "Supervisors", "Internal: Ms. Nimali Hettiarachchi & Ms. Geethika Senarathne | External: Lead Software Engineer", 13, dark_text=False)

    set_speaker_notes(s1, 
        "Welcome everyone. Good morning respected lecturers and evaluation panel members. "
        "Today, I am proud to present my Industrial Training project: 'KNW PRO Realtor' — an enterprise-grade digital property trading, "
        "title deed verification, and legal advisory platform developed during my six-month internship at KNW PRO Realtor (Pvt) Ltd in Colombo. "
        "In this presentation, I will walk you through the real-world real estate problems we identified, the decoupled architecture we engineered, "
        "our key innovation in unified single-account dynamic role switching, and how we bridged software engineering with legal deed security."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 2: Presentation Agenda
    # ═════════════════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, COLOR_BG_LIGHT)
    add_header(s2, "Overview", "Presentation Roadmap & Agenda")

    # 4 Agenda Cards
    agenda_items = [
        ("01. Context & Motivation", [
            "Host Organization: KNW PRO Realtor",
            "Sri Lankan Property Market Challenges",
            "Title Deed Insecurity & Fraud Risks",
            "Project Objectives & Value Proposition"
        ]),
        ("02. Architectural Design", [
            "Decoupled 3-Tier Architecture",
            "Frontend: React 18, Vite & Context APIs",
            "Backend: PHP 8.4 & Laravel 13 REST API",
            "Database: MySQL 3NF Relational Model"
        ]),
        ("03. Key Innovations & Modules", [
            "Unified Single-Account Dynamic Role Switcher",
            "5-Step Property Intake & Pricing Estimator",
            "LegalTech Deed Review & Lawyer Scheduling",
            "Admin Moderation & PayHere Integration"
        ]),
        ("04. Engineering & Evaluation", [
            "Agile Scrum & GitFlow Collaboration",
            "Multi-Tier Testing & Security Hardening",
            "Technical Challenges & Solutions",
            "Competencies Gained & Future Scalability"
        ]),
    ]

    for idx, (title, points) in enumerate(agenda_items):
        col = idx % 2
        row = idx // 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.8 + row * 2.6)
        _, card_tf = add_card(s2, left, top, Inches(5.7), Inches(2.4), title, COLOR_EMERALD)
        for pt in points:
            add_bullet_point(card_tf, "", pt, 11)

    set_speaker_notes(s2,
        "Here is the roadmap for today's presentation. I have structured our walkthrough into four clear segments: "
        "First, the background of our host enterprise and the critical industry problems that motivated this platform. "
        "Second, the full-stack architecture connecting React, Laravel, and MySQL. "
        "Third, our major functional achievements, highlighting our unified account role-switching engine and legal deed verification conduits. "
        "And finally, our quality assurance practices, solutions to difficult technical challenges, and my professional growth during the training."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 3: Host Organization Overview
    # ═════════════════════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, COLOR_BG_LIGHT)
    add_header(s3, "Host Enterprise", "KNW PRO Realtor (Private) Limited")

    # Left Column: Company Profile & Business Pillars
    _, left_tf = add_card(s3, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), "Corporate Profile & Operations", COLOR_EMERALD)
    add_bullet_point(left_tf, "Enterprise Identity", "Premier real estate brokerage, property consultancy, and land sales enterprise headquartered in Colombo, Sri Lanka.")
    add_bullet_point(left_tf, "Market Reach", "Active across Western, Central, and Southern Provinces managing residential homes, commercial complexes, and subdivided land plots.")
    add_bullet_point(left_tf, "Core Services", "Real estate sales & acquisitions, land demarcation & appraisal, legal deed verification coordination, and property investment consulting.")
    add_bullet_point(left_tf, "Digital Modernization", "Established an in-house IT & Digital Systems Department to pioneer secure, fraud-resilient property tech solutions.")

    # Right Column: Department of Study & Internship Placement
    _, right_tf = add_card(s3, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "IT & Digital Systems Division (Placement)", COLOR_GOLD)
    add_bullet_point(right_tf, "Department Mission", "Engineering proprietary cloud platforms, internal databases, digital document vaults, and client portals.")
    add_bullet_point(right_tf, "Trainee's Role", "Software Engineering Intern actively embedded in full-stack feature development, API architecture, and database design.")
    add_bullet_point(right_tf, "Working Culture", "Agile Scrum framework, 2-week sprint cycles, daily standups (09:30 AM), GitFlow peer code reviews.")
    add_bullet_point(right_tf, "Interdisciplinary Exposure", "Direct collaborative engagement with senior developers, licensed conveyancing attorneys, and real estate brokers.")

    set_speaker_notes(s3,
        "Let me introduce my host organization. KNW PRO Realtor (Private) Limited is a leading real estate brokerage and land consultancy enterprise in Sri Lanka. "
        "Unlike conventional agencies that rely solely on offline brokers and paper documents, KNW PRO Realtor made the strategic decision to build an in-house "
        "IT & Digital Systems Department. My internship placement was within this engineering team. Our daily culture followed Agile Scrum with two-week sprints, "
        "peer code reviews, and direct collaboration with licensed attorneys and property valuation consultants."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 4: Problem Statement & Motivation
    # ═════════════════════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, COLOR_BG_LIGHT)
    add_header(s4, "Problem Statement", "Critical Bottlenecks in the Sri Lankan Property Market")

    problems = [
        ("1. Title Deed Fraud & Bogus Listings", 
         "Unregistered middlemen and unauthorized third parties advertise land they do not own. Buyers risk losing life savings on properties burdened by undisclosed encumbrances, fake deeds, or local council boundary disputes.",
         RGBColor(239, 68, 68)), # Red
        ("2. Total Absence of Pre-Listing Legal Screening", 
         "Conventional property websites merely publish unvetted photos and phone numbers. Prospective buyers discover major legal defects only after paying earnest money deposits and hiring private lawyers.",
         COLOR_GOLD),
        ("3. Rigid Multi-Account Registration Friction", 
         "Existing software architectures assign a single immutable role per email. When an active property buyer wants to sell a house, the system blocks them, demanding logout and registration with a completely separate second email.",
         COLOR_BLUE),
    ]

    for idx, (p_title, p_desc, border_c) in enumerate(problems):
        top = Inches(1.8 + idx * 1.7)
        _, p_tf = add_card(s4, Inches(0.8), top, Inches(11.7), Inches(1.5), p_title, border_c)
        add_bullet_point(p_tf, "Impact & Challenge", p_desc, 11)

    set_speaker_notes(s4,
        "Why did we build KNW PRO Realtor? The Sri Lankan real estate industry suffers from three severe problems: "
        "First, rampant title deed fraud. Bogus listings are uploaded by unauthorized middlemen, resulting in protracted legal battles. "
        "Second, existing listing portals do not perform any pre-listing legal screening; buyers discover forged deeds only after placing deposits. "
        "And third, from a software UX standpoint, legacy platforms force users to maintain separate accounts with multiple email addresses if they want to both buy and sell property. "
        "These three challenges shaped the core engineering requirements of our project."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 5: Proposed Solution — KNW PRO Realtor
    # ═════════════════════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, COLOR_BG_LIGHT)
    add_header(s5, "The Solution", "KNW PRO Realtor: Digital Property & Legal Advisory Platform")

    features = [
        ("Unified Single Identity", "One user account for life. Dynamic role switching between Buyer, Seller, and Lawyer modes with in-place NIC verification and zero session loss.", COLOR_EMERALD),
        ("Legal Verification Guarantee", "Direct integration with licensed attorneys-at-law who audit Land Registry extracts, survey plans, and assign verified certification badges.", COLOR_GOLD),
        ("Multi-Step Listing Intake", "Comprehensive 5-step property wizard capturing location coordinates, specs, high-res photos, title deeds, and automated price valuation estimation.", COLOR_BLUE),
        ("Buyer Engagement Hub", "Personalized property bookmarks, custom user inspection notes, one-click printable property dossiers, and direct physical site visit booking.", COLOR_EMERALD),
        ("Legal Consultation Portal", "Browse verified property attorneys, select calendar slots, book video consultations, and download RFC-compliant iCalendar (.ics) invite files.", COLOR_GOLD),
        ("Admin Governance & Auditing", "Centralized moderation dashboard for National Identity Card approval, payment receipt reconciliation, and listing quality control.", COLOR_BLUE),
    ]

    for idx, (title, desc, color) in enumerate(features):
        col = idx % 3
        row = idx // 3
        left = Inches(0.8 + col * 4.0)
        top = Inches(1.8 + row * 2.6)
        _, f_tf = add_card(s5, left, top, Inches(3.7), Inches(2.4), title, color)
        add_bullet_point(f_tf, "", desc, 10.5)

    set_speaker_notes(s5,
        "Our solution is KNW PRO Realtor — a modern, secure web platform designed around six core pillars: "
        "A Unified Single Identity that allows users to seamlessly switch between buying and selling without juggling duplicate accounts. "
        "A Legal Verification Guarantee that integrates certified attorneys directly into the listing pipeline. "
        "A rich multi-step intake wizard with valuation calculators. "
        "A dedicated Buyer Engagement Hub with saved shortlists and printable dossiers. "
        "A LegalTech portal with automated iCalendar scheduling. "
        "And an Administrative Governance Dashboard ensuring quality and regulatory compliance."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 6: System Architecture & Decoupled Model
    # ═════════════════════════════════════════════════════════════════════════
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, COLOR_BG_LIGHT)
    add_header(s6, "System Architecture", "Decoupled 3-Tier Client-Server Architecture")

    # 3 Architecture Tier Cards
    tiers = [
        ("Tier 1: Client Application (SPA)", [
            "Technology: React 18, Vite Bundler, Modern CSS3",
            "Single Page Application with fast client-side routing",
            "Context State: AuthContext, LanguageContext, CurrencyContext",
            "Asynchronous HTTP communication via Axios with interceptors",
            "Responsive layout optimized for Mobile, Tablet, and Desktop"
        ], COLOR_BLUE),
        ("Tier 2: Business Logic & REST API", [
            "Technology: PHP 8.4 on Laravel 13 Framework",
            "Model-View-Controller (MVC) architectural pattern",
            "RESTful JSON endpoints with FormRequest validation",
            "Bcrypt cryptographic hashing (12 rounds) & RBAC middleware",
            "Private disk storage isolation for sensitive deeds & NIC files"
        ], COLOR_EMERALD),
        ("Tier 3: Relational Persistence", [
            "Technology: MySQL 8.0 Relational Database",
            "Normalized up to Third Normal Form (3NF)",
            "Eloquent ORM abstraction with strongly typed models",
            "Version-controlled schema blueprints via Laravel Migrations",
            "Composite indexing on district, price, and publication status"
        ], COLOR_GOLD),
    ]

    for idx, (title, points, color) in enumerate(tiers):
        left = Inches(0.8 + idx * 4.0)
        _, t_tf = add_card(s6, left, Inches(1.8), Inches(3.7), Inches(5.1), title, color)
        for pt in points:
            add_bullet_point(t_tf, "", pt, 11)

    set_speaker_notes(s6,
        "Here we examine the architectural foundation of KNW PRO Realtor. We selected a decoupled three-tier client-server architecture. "
        "Tier 1 is the Single Page Application built with React 18 and Vite, delivering instantaneous client-side transitions and reactive global state. "
        "Tier 2 is our server-side API constructed with PHP 8.4 on Laravel 13, adhering strictly to RESTful standards, FormRequest validation, and private disk isolation. "
        "Tier 3 is our relational MySQL database, strictly normalized to Third Normal Form to ensure zero data redundancy and ACID transaction reliability."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 7: Frontend Architecture & Global State
    # ═════════════════════════════════════════════════════════════════════════
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, COLOR_BG_LIGHT)
    add_header(s7, "Frontend Architecture", "React 18 Component Hierarchy & Global Context State")

    _, f_left_tf = add_card(s7, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), "Reactive Context Providers", COLOR_BLUE)
    add_bullet_point(f_left_tf, "AuthContext", "Manages user identity, bearer tokens, capability claims (can.act_as_buyer, can.act_as_seller), and activeRole state switching.")
    add_bullet_point(f_left_tf, "LanguageContext", "Trilingual internationalization (i18n) supporting English, Sinhala, and Tamil translations across all interface labels.")
    add_bullet_point(f_left_tf, "CurrencyContext", "Real-time exchange conversion between Sri Lankan Rupees (LKR), USD, EUR, and GBP with localized symbol formatting.")
    add_bullet_point(f_left_tf, "Axios Interceptors", "Centralized request header token attachment and global HTTP 401/403/422 error response interceptors.")

    _, f_right_tf = add_card(s7, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "Component-Driven UI Engineering", COLOR_EMERALD)
    add_bullet_point(f_right_tf, "Modular Components", "Over 25 reusable functional components: Navbar, PropertyCard, Modal, PriceEstimator, ScheduleVisitModal.")
    add_bullet_point(f_right_tf, "Debounced Live Search", "Implemented 400ms debounce hook on marketplace search inputs, preventing excessive API roundtrips.")
    add_bullet_point(f_right_tf, "Form State Recovery", "Lifting multi-step draft wizard state into centralized reducers with sessionStorage persistence.")
    add_bullet_point(f_right_tf, "CSS3 Design System", "Strict design tokens using CSS variables, flexbox, and CSS grid layouts ensuring pixel-perfect responsiveness.")

    set_speaker_notes(s7,
        "Moving to the frontend: our React application eliminates prop-drilling by establishing three dedicated global context providers. "
        "AuthContext maintains authentication and dynamic role states. LanguageContext provides instant trilingual translation between English, Sinhala, and Tamil. "
        "And CurrencyContext dynamically recalculates property prices across LKR, USD, EUR, and GBP. "
        "We also implemented input debouncing to prevent server overload and centralized form reducers so sellers never lose entered data when navigating wizard steps."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 8: Backend Architecture & REST API Design
    # ═════════════════════════════════════════════════════════════════════════
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, COLOR_BG_LIGHT)
    add_header(s8, "Backend Architecture", "Laravel 13 REST API & Eloquent ORM Pipeline")

    _, b_left_tf = add_card(s8, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), "RESTful Controllers & Routes", COLOR_EMERALD)
    add_bullet_point(b_left_tf, "AuthController", "Registration, credential authentication, user resource formatting, and in-place applyForSeller verification.")
    add_bullet_point(b_left_tf, "SellerPropertyDraftController", "Multi-step draft intake, gallery upload streaming, and listing publication checks via $user->canListProperty().")
    add_bullet_point(b_left_tf, "LawyerController", "Attorney profile management, weekly availability slot configuration, appointment booking, and deed review audits.")
    add_bullet_point(b_left_tf, "AdminUserController", "Identity verification approval/rejection, property quality moderation, and system revenue analytics.")

    _, b_right_tf = add_card(s8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "Database & Security Pipeline", COLOR_GOLD)
    add_bullet_point(b_right_tf, "Eloquent ORM", "Declarative relationships (hasMany, belongsTo) with eager-loading scopes to eliminate N+1 query performance hits.")
    add_bullet_point(b_right_tf, "Idempotent Migrations", "Version-controlled schema blueprints enabling repeatable, automated database provisioning across environments.")
    add_bullet_point(b_right_tf, "Security Hardening", "Bcrypt hashing with 12 rounds, CSRF token validation, and MIME-type validation on all uploaded files.")
    add_bullet_point(b_right_tf, "Private Storage Disk", "Sensitive NIC documents and original title deeds are isolated from public web directories.")

    set_speaker_notes(s8,
        "On the backend, Laravel 13 organizes business logic across dedicated RESTful controllers. "
        "We use FormRequests to enforce validation before controller execution, ensuring data integrity. "
        "Eloquent ORM manages complex database relationships with eager-loading to prevent N+1 query penalties. "
        "From a security perspective, all passwords use 12-round Bcrypt hashing, and sensitive documents like National Identity Cards are stored in a private disk "
        "accessible only through authenticated administrative proxy routes."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 9: Major Innovation — Unified Role-Switching
    # ═════════════════════════════════════════════════════════════════════════
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, COLOR_BG_LIGHT)
    add_header(s9, "Core Innovation", "Unified Single-Account Dynamic Role-Switching Engine")

    # Before vs After Comparison
    _, left_cmp_tf = add_card(s9, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), "Legacy Multi-Account Flaw", RGBColor(239, 68, 68))
    add_bullet_point(left_cmp_tf, "Rigid Role Column", "Users table held a static role enum ('buyer', 'seller', 'lawyer'). One email was permanently restricted to one role.")
    add_bullet_point(left_cmp_tf, "Registration Failure", "When a buyer clicked 'Sell', system redirected to /register?role=seller, which failed due to duplicate email uniqueness.")
    add_bullet_point(left_cmp_tf, "Lawyer Isolation", "Attorneys were locked inside LawyerDashboard and could not browse the marketplace or bookmark properties.")
    add_bullet_point(left_cmp_tf, "Poor User Experience", "Users were forced to logout and login repeatedly using multiple different email credentials.")

    _, right_cmp_tf = add_card(s9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "Newly Engineered Unified Architecture", COLOR_EMERALD)
    add_bullet_point(right_cmp_tf, "Decoupled Seller Status", "Added seller_status enum ('none', 'pending', 'active', 'rejected'), decoupling identity from selling rights.")
    add_bullet_point(right_cmp_tf, "In-Place Verification", "Existing buyers submit NIC verification via POST /api/user/apply-seller directly from their active session.")
    add_bullet_point(right_cmp_tf, "Dynamic Capabilities", "User resource returns boolean capability flags: act_as_buyer, act_as_seller, is_pending_seller, is_lawyer.")
    add_bullet_point(right_cmp_tf, "1-Click Navbar Switcher", "Interactive mode switcher pills in Navbar allow instant toggling between Buyer Mode, Seller Mode, and Lawyer Mode.")

    set_speaker_notes(s9,
        "This slide highlights our most significant architectural innovation. In legacy systems, users were trapped in a single role. "
        "If a property buyer wanted to list a house, the system demanded they register a second account with a second email address, causing immense user drop-off. "
        "We completely re-architected the identity model. We added a decoupled `seller_status` column and built an in-place verification API. "
        "Now, every user holds a single credential. Unverified buyers can submit their NIC directly from the Sell page, and once approved, toggle instantly "
        "between Buyer Mode, Seller Mode, and Lawyer Mode with a single click in the navigation bar."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 10: LegalTech & Deed Verification Workflow
    # ═════════════════════════════════════════════════════════════════════════
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, COLOR_BG_LIGHT)
    add_header(s10, "LegalTech Integration", "Title Deed Verification & Lawyer Consultation System")

    steps = [
        ("Step 1: Document Submission", "Seller uploads scanned property title deed, certified survey plan, and local authority street line clearance certificate during intake wizard.", COLOR_BLUE),
        ("Step 2: Attorney Deed Audit", "Uploaded legal dossiers are routed to certified real estate attorneys-at-law who conduct Land Registry volume/folio extract verification.", COLOR_GOLD),
        ("Step 3: Verification Badging", "Cleared properties earn the official 'KNW Certified Deed Verified' badge, signaling 100% legal title safety to prospective buyers.", COLOR_EMERALD),
        ("Step 4: Consultation Booking", "Buyers can book direct video consultations with real estate lawyers to discuss deeds, boundary lines, and purchase agreements.", COLOR_BLUE),
        ("Step 5: Automated Scheduling", "Platform generates RFC 5545 compliant iCalendar (.ics) invitations and dispatches automated email confirmations to both parties.", COLOR_GOLD),
        ("Step 6: Audit Trail Logging", "Full legal audit history is permanently recorded in the database, ensuring complete transaction accountability.", COLOR_EMERALD),
    ]

    for idx, (title, desc, color) in enumerate(steps):
        col = idx % 3
        row = idx // 3
        left = Inches(0.8 + col * 4.0)
        top = Inches(1.8 + row * 2.6)
        _, s_tf = add_card(s10, left, top, Inches(3.7), Inches(2.4), title, color)
        add_bullet_point(s_tf, "", desc, 10.5)

    set_speaker_notes(s10,
        "Here we see the six-step LegalTech verification workflow that protects buyers from fraud. "
        "When a seller lists a property, deeds and survey plans are submitted into our secure vault. "
        "Licensed attorneys review Land Registry extracts before approving the listing. Only cleared properties receive the 'KNW Certified Deed Verified' badge. "
        "Furthermore, buyers can schedule video consultations with property attorneys, receiving automated iCalendar invites directly into their Google or Outlook calendars."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 11: Monetization, Inquiries & Admin Governance
    # ═════════════════════════════════════════════════════════════════════════
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, COLOR_BG_LIGHT)
    add_header(s11, "Commercial Operations", "Monetization Tiers, Inquiries & Administrative Governance")

    _, m_left_tf = add_card(s11, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), "Monetization & Inquiries", COLOR_GOLD)
    add_bullet_point(m_left_tf, "Listing Tiers", "Basic (free 30-day listing), Verified (attorney deed review included), and Featured (homepage carousel & priority search placement).")
    add_bullet_point(m_left_tf, "PayHere Gateway", "Automated checkout with MD5 signature validation for instant card payments and merchant IPN callback processing.")
    add_bullet_point(m_left_tf, "Bank Deposit Fallback", "Allows traditional sellers to upload bank deposit slips for manual administrative reconciliation.")
    add_bullet_point(m_left_tf, "Site Visit Inquiries", "Buyers request physical inspection dates; sellers receive instant notifications to accept or propose alternative slots.")

    _, m_right_tf = add_card(s11, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "Admin Operations & Moderation", COLOR_BLUE)
    add_bullet_point(m_right_tf, "Seller Identity Queue", "Inspects uploaded NIC/Passport documents in private storage; approves verified sellers or records structured rejection reasons.")
    add_bullet_point(m_right_tf, "Draft Quality Moderation", "Screens property photos, validates price reasonableness, and toggles published visibility on the marketplace.")
    add_bullet_point(m_right_tf, "Payment Verification", "Audits uploaded bank deposit receipts against banking records before activating premium tier badges.")
    add_bullet_point(m_right_tf, "Executive KPI Dashboard", "Real-time analytics tracking registered buyers, active sellers, approved listings, and platform revenue.")

    set_speaker_notes(s11,
        "Slide 11 demonstrates the commercial and administrative operations of the platform. "
        "We implemented a three-tier monetization model with PayHere payment gateway integration and bank deposit slip fallback. "
        "For buyers, the system facilitates site visit requests with calendar synchronization. "
        "And for enterprise operations, the Admin Dashboard provides centralized control: moderating seller identity applications, auditing property drafts, and tracking business KPIs."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 12: SDLC, Testing & Security Hardening
    # ═════════════════════════════════════════════════════════════════════════
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, COLOR_BG_LIGHT)
    add_header(s12, "Engineering Quality", "Software Development Life Cycle, Testing & Security")

    _, q_left_tf = add_card(s12, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), "Agile SDLC & GitFlow", COLOR_EMERALD)
    add_bullet_point(q_left_tf, "Scrum Methodology", "Executed 2-week sprint cycles with daily standups, backlog refinement, sprint demos, and retrospectives.")
    add_bullet_point(q_left_tf, "GitFlow Branching", "Main, develop, and feature-specific branches (feature/role-switching, feature/payhere-ipn) with mandatory pull request reviews.")
    add_bullet_point(q_left_tf, "Unit & Integration Testing", "PHPUnit test cases on backend controller validation and model relationships; automated Postman regression suites.")
    add_bullet_point(q_left_tf, "Production Build", "Vite production bundling completed in 5.16s with asset minification, chunk splitting, and zero linting errors.")

    _, q_right_tf = add_card(s12, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "Security & Compliance Hardening", COLOR_GOLD)
    add_bullet_point(q_right_tf, "Bcrypt Cryptographic Hashing", "Passwords hashed with 12 rounds, resilient against rainbow table and dictionary attacks.")
    add_bullet_point(q_right_tf, "Private Disk Document Vault", "Sensitive NIC documents and original title deeds are isolated from public web directories.")
    add_bullet_point(q_right_tf, "SQLi & XSS Prevention", "Eloquent PDO parameter binding eliminates SQL injection; React virtual DOM auto-escaping neutralizes XSS threats.")
    add_bullet_point(q_right_tf, "Cross-Browser Verification", "Tested across Chrome, Firefox, Edge, and Safari on iOS, Android, and desktop viewports.")

    set_speaker_notes(s12,
        "Software engineering rigor was maintained throughout the development lifecycle. "
        "We followed Agile Scrum in two-week sprints and governed code quality with GitFlow branching and peer reviews. "
        "Our testing strategy encompassed PHPUnit unit tests, Postman API collections, and cross-browser responsive audits. "
        "Security was prioritized from day one: Bcrypt password hashing, private storage isolation for identity documents, and complete protection against SQL injection and XSS."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 13: Technical Challenges & Solutions
    # ═════════════════════════════════════════════════════════════════════════
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13, COLOR_BG_LIGHT)
    add_header(s13, "Problem Solving", "Key Technical Challenges & Implemented Solutions")

    challenges = [
        ("Duplicate Account Collision", 
         "Database-level uniqueness on user email prevented buyers from creating seller accounts.",
         "Decoupled seller verification into a seller_status column and built in-place applyForSeller endpoint, enabling a single credential to hold multiple roles.",
         COLOR_EMERALD),
        ("Lawyer Cross-Role Isolation", 
         "Lawyers were locked in LawyerDashboard with hidden navbars, unable to browse or save properties.",
         "Engineered bidirectional cross-navigation in Navbar and dashboards, allowing attorneys to seamlessly alternate between legal and buyer modes.",
         COLOR_BLUE),
        ("Title Deed & Identity Exposure", 
         "Direct public URLs to uploaded deeds and NIC cards created severe data privacy vulnerabilities.",
         "Stored all sensitive files in non-public storage directories (storage/app/private) accessible solely via authenticated admin proxy routes.",
         COLOR_GOLD),
        ("Client Session Desynchronization", 
         "Role changes caused stale data in browser localStorage, requiring hard page reloads to reflect permissions.",
         "Implemented refreshUser() dispatcher in AuthContext to synchronize React memory state with localStorage instantly upon role transitions.",
         COLOR_EMERALD),
    ]

    for idx, (c_name, c_prob, c_sol, color) in enumerate(challenges):
        col = idx % 2
        row = idx // 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.8 + row * 2.6)
        _, c_tf = add_card(s13, left, top, Inches(5.7), Inches(2.4), c_name, color)
        add_bullet_point(c_tf, "Root Cause", c_prob, 10.5)
        add_bullet_point(c_tf, "Solution", c_sol, 10.5)

    set_speaker_notes(s13,
        "Every complex software project presents technical hurdles. This slide summarizes four major challenges we overcame: "
        "First, resolving the duplicate email collision by decoupling selling authorization from identity. "
        "Second, breaking the isolation of lawyer accounts by introducing bidirectional cross-navigation. "
        "Third, eliminating document exposure risks by enforcing private disk isolation and authenticated streaming. "
        "And fourth, resolving client-side session desynchronization with reactive state refresh dispatchers."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 14: Competencies & Future Scalability
    # ═════════════════════════════════════════════════════════════════════════
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14, COLOR_BG_LIGHT)
    add_header(s14, "Learning & Evolution", "Professional Growth & Future Scalability Roadmap")

    _, l_tf = add_card(s14, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), "Competencies & Skills Acquired", COLOR_EMERALD)
    add_bullet_point(l_tf, "Full-Stack Proficiency", "Advanced practical expertise in PHP 8.4, Laravel 13, React 18, Vite, and MySQL relational modeling.")
    add_bullet_point(l_tf, "Architectural Problem-Solving", "Ability to diagnose architectural friction points and implement clean, backward-compatible refactors.")
    add_bullet_point(l_tf, "Enterprise Security", "Mastered Bcrypt hashing, private disk isolation, FormRequest validation, and RBAC authorization middleware.")
    add_bullet_point(l_tf, "Professional Collaboration", "Proficiency in Agile Scrum rituals, GitFlow version control, sprint planning, and cross-disciplinary communication.")

    _, r_tf = add_card(s14, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "Future Scalability Recommendations", COLOR_GOLD)
    add_bullet_point(r_tf, "Automated OCR Identity Extraction", "Integrate Optical Character Recognition (OCR) to automatically parse NIC numbers and photo matches from uploads.")
    add_bullet_point(r_tf, "Bidirectional WebSockets", "Upgrade inquiry notifications and appointment alerts to real-time WebSockets using Laravel Reverb.")
    add_bullet_point(r_tf, "Cross-Platform Mobile Apps", "Develop iOS and Android mobile applications using React Native sharing the existing Laravel REST backend.")
    add_bullet_point(r_tf, "Geospatial Boundary Mapping", "Integrate Leaflet or Google Maps APIs for cadastral property boundary and survey plan overlays.")

    set_speaker_notes(s14,
        "Reflecting on this training experience: I gained profound growth across full-stack software development, architectural refactoring, and security best practices. "
        "Looking forward, we have identified several exciting avenues for future scalability: "
        "Automated OCR for scanning identity cards, real-time push notifications using Laravel Reverb WebSockets, cross-platform mobile apps using React Native, "
        "and interactive GIS mapping for property boundaries."
    )

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 15: Conclusion & Acknowledgements (Dark Navy)
    # ═════════════════════════════════════════════════════════════════════════
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_background(s15, COLOR_BG_DARK)

    # Decorative bar
    bar15 = s15.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    bar15.fill.solid()
    bar15.fill.fore_color.rgb = COLOR_GOLD
    bar15.line.fill.background()

    # Heading
    concl_box = s15.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(1.5))
    c_tf = concl_box.text_frame
    c_tf.word_wrap = True
    cp1 = c_tf.paragraphs[0]
    cp1.text = "CONCLUSION & ACKNOWLEDGEMENTS"
    cp1.font.name = FONT_TITLE
    cp1.font.size = Pt(32)
    cp1.font.bold = True
    cp1.font.color.rgb = COLOR_WHITE
    cp1.space_after = Pt(8)

    cp2 = c_tf.add_paragraph()
    cp2.text = "Module IT3901 Industrial Training Successfully Completed"
    cp2.font.name = FONT_TITLE
    cp2.font.size = Pt(18)
    cp2.font.color.rgb = COLOR_GOLD

    # Two Summary Cards
    _, c_left_tf = add_card(s15, Inches(1.0), Inches(2.8), Inches(5.5), Inches(3.8), "Project Achievements", COLOR_GOLD, COLOR_CARD_DARK)
    add_bullet_point(c_left_tf, "Production-Ready Platform", "Successfully built and deployed the full-stack KNW PRO Realtor real estate and legal clearance platform.", 12, dark_text=False)
    add_bullet_point(c_left_tf, "Academic Fulfillment", "100% fulfilled the industrial training curriculum requirements for the Higher Diploma in Software Engineering at LNBTI.", 12, dark_text=False)
    add_bullet_point(c_left_tf, "Commercial Impact", "Delivered a fraud-resilient property ecosystem with seamless single-account role switching.", 12, dark_text=False)

    _, c_right_tf = add_card(s15, Inches(6.8), Inches(2.8), Inches(5.5), Inches(3.8), "Sincere Acknowledgements", COLOR_EMERALD, COLOR_CARD_DARK)
    add_bullet_point(c_right_tf, "Academic Mentors", "Profound gratitude to Lanka Nippon BizTech Institute, faculty members, and internal supervisors.", 12, dark_text=False)
    add_bullet_point(c_right_tf, "Host Organization", "Heartfelt thanks to KNW PRO Realtor (Private) Limited management, external supervisors, and engineering colleagues.", 12, dark_text=False)
    add_bullet_point(c_right_tf, "Open for Questions", "Thank you for your time, attention, and guidance. I welcome questions from the evaluation panel.", 12, dark_text=False)

    set_speaker_notes(s15,
        "In conclusion, the six-month industrial training program at KNW PRO Realtor (Private) Limited has been an invaluable, transformative milestone in my engineering education. "
        "The project successfully delivered a robust, secure, and production-ready platform that solves genuine commercial real estate problems in Sri Lanka. "
        "I would like to express my deepest gratitude to my academic supervisors at LNBTI, my technical mentors at KNW PRO Realtor, and my family for their unwavering support. "
        "Thank you very much. I am now open to answer any questions from the panel."
    )

    output_path = os.path.abspath("KNW_PRO_Realtor_Industrial_Project_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation successfully generated at: {output_path}")

if __name__ == "__main__":
    build_presentation()
