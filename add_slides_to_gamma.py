import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PPTX_PATH = "KNW-PRO-Realtor.pptx"

COLOR_TAG = RGBColor(125, 110, 96)       # #7D6E60 Warm Slate
COLOR_TITLE = RGBColor(42, 37, 32)       # #2A2520 Charcoal/Black
COLOR_BODY = RGBColor(107, 93, 79)       # #6B5D4F Muted text
COLOR_BORDER = RGBColor(226, 232, 240)   # #E2E8F0 Subtle Border
COLOR_CARD_BG = RGBColor(255, 255, 255)  # White Card
COLOR_CARD_TINT = RGBColor(248, 250, 252)# Soft grey tint
COLOR_PRIMARY = RGBColor(37, 99, 235)    # Blue accent
COLOR_EMERALD = RGBColor(16, 185, 129)   # Emerald accent
COLOR_AMBER = RGBColor(217, 119, 6)      # Amber accent

FONT_HEADING = "Arial"
FONT_BODY = "Arial"

def add_header_and_footer(slide, tag_text, title_text, slide_number):
    # Category Tag / Subtitle
    tag_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.33), Inches(0.3))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.name = FONT_BODY
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_TAG

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.7), Inches(12.33), Inches(0.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TITLE

    # Footer Line
    f_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(7.0), Inches(12.33), Inches(0.01))
    f_line.fill.solid()
    f_line.fill.fore_color.rgb = COLOR_BORDER
    f_line.line.color.rgb = COLOR_BORDER

    # Footer Left
    f_left = slide.shapes.add_textbox(Inches(0.5), Inches(7.08), Inches(4.0), Inches(0.3))
    tf_fleft = f_left.text_frame
    tf_fleft.margin_left = tf_fleft.margin_top = tf_fleft.margin_right = tf_fleft.margin_bottom = 0
    p_fl = tf_fleft.paragraphs[0]
    p_fl.text = "KNW PRO Realtor • System Architecture & Engineering"
    p_fl.font.name = FONT_BODY
    p_fl.font.size = Pt(9.5)
    p_fl.font.color.rgb = COLOR_BODY

    # Footer Right
    f_right = slide.shapes.add_textbox(Inches(10.83), Inches(7.08), Inches(2.0), Inches(0.3))
    tf_fright = f_right.text_frame
    tf_fright.margin_left = tf_fright.margin_top = tf_fright.margin_right = tf_fright.margin_bottom = 0
    p_fr = tf_fright.paragraphs[0]
    p_fr.alignment = PP_ALIGN.RIGHT
    p_fr.text = f"{slide_number}"
    p_fr.font.name = FONT_BODY
    p_fr.font.size = Pt(9.5)
    p_fr.font.color.rgb = COLOR_BODY

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card

prs = Presentation(PPTX_PATH)
blank_layout = prs.slide_layouts[0]

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 11: FUNCTIONAL REQUIREMENTS BREAKDOWN (4 USERS)
# ─────────────────────────────────────────────────────────────────────────────
s11 = prs.slides.add_slide(blank_layout)
add_header_and_footer(s11, "REQUIREMENTS SPECIFICATION", "Detailed Functional Requirements by Stakeholder Role", 11)

col_w = Inches(2.88)
card_h = Inches(5.1)
gap = Inches(0.24)
start_x = Inches(0.5)
card_y = Inches(1.5)

roles_data = [
    ("Property Buyer", COLOR_PRIMARY, [
        ("FR-B1 Property Discovery", "District, property type & price filtering with instant search index."),
        ("FR-B2 Deed Verification", "Check 'KNW Certified Deed Verified' badge and legal summary dossier."),
        ("FR-B3 Financial Tools", "Mortgage amortization & stamp duty payment calculations."),
        ("FR-B4 Physical Site Visits", "Schedule in-person inspection dates and track confirmation."),
        ("FR-B5 Legal Consultation", "Browse attorneys, view available calendar slots & book video advice.")
    ]),
    ("Property Seller", COLOR_EMERALD, [
        ("FR-S1 In-Place Verification", "Submit Phone, Address & NIC document without re-registering."),
        ("FR-S2 Listing Creation", "Multi-step draft intake capturing coordinates, features & price."),
        ("FR-S3 Legal Uploads", "Securely upload Title Deeds, Survey Plans & Local Council docs."),
        ("FR-S4 Monetization Tiers", "Basic, Verified, and Featured listing plans via PayHere gateway."),
        ("FR-S5 Inquiry Handling", "Respond to buyer inquiries and manage site visit bookings.")
    ]),
    ("Real Estate Lawyer", COLOR_AMBER, [
        ("FR-L1 Profile & Fees", "Manage bar credentials, hourly rates, bio, and languages."),
        ("FR-L2 Availability Slots", "Define weekly consultation schedule & block holiday dates."),
        ("FR-L3 Deed Review Queue", "Access assigned property title deeds and survey cadastrals."),
        ("FR-L4 Legal Clearance", "Issue deed review findings and recommend verified badge."),
        ("FR-L5 Video Consultations", "Host video advisory meetings with calendar invite sync.")
    ]),
    ("System Administrator", COLOR_TAG, [
        ("FR-A1 Identity Governance", "Audit pending seller NIC submissions; approve or reject."),
        ("FR-A2 Listing Quality Audit", "Screen property drafts and validate payment receipts/slips."),
        ("FR-A3 Deed Badge Control", "Toggle official 'KNW Certified Deed Verified' badge."),
        ("FR-A4 Payment Audit", "Oversee listing fee & lawyer consultation fee transactions."),
        ("FR-A5 Platform Analytics", "Real-time metrics: active users, revenue, and pending tasks.")
    ])
]

for idx, (role_name, accent, items) in enumerate(roles_data):
    cx = start_x + idx * (col_w + gap)
    add_card(s11, cx, card_y, col_w, card_h, COLOR_CARD_BG, COLOR_BORDER)
    
    # Accent top pill
    pill = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, card_y, col_w, Inches(0.08))
    pill.fill.solid()
    pill.fill.fore_color.rgb = accent
    pill.line.fill.background()
    
    # Text Box
    tbox = s11.shapes.add_textbox(cx + Inches(0.15), card_y + Inches(0.18), col_w - Inches(0.3), card_h - Inches(0.3))
    tf = tbox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p0 = tf.paragraphs[0]
    p0.text = role_name
    p0.font.name = FONT_HEADING
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_TITLE
    p0.space_after = Pt(10)
    
    for title, desc in items:
        p_t = tf.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = accent
        
        p_d = tf.add_paragraph()
        p_d.text = f"  {desc}"
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_BODY
        p_d.space_after = Pt(6)

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 12: UNIFIED SYSTEM USE CASE ARCHITECTURE
# ─────────────────────────────────────────────────────────────────────────────
s12 = prs.slides.add_slide(blank_layout)
add_header_and_footer(s12, "USE CASE MODELING", "Unified System Use Case Architecture (All 4 Actors)", 12)

# Left Column: Image Card
img_card_w = Inches(5.0)
img_card_h = Inches(5.1)
add_card(s12, Inches(0.5), Inches(1.5), img_card_w, img_card_h, COLOR_CARD_BG, COLOR_BORDER)
if os.path.exists("diagrams/uc_unified.png"):
    s12.shapes.add_picture("diagrams/uc_unified.png", Inches(1.5), Inches(1.6), width=Inches(3.0))

# Right Column: Explanatory Architecture Cards
info_w = Inches(7.05)
info_x = Inches(5.78)

arch_points = [
    ("Unified Identity & Dynamic Role Switching", "Eliminated legacy multi-account friction. Every user defaults to Property Buyer. Users activate Seller privileges in-place via NIC upload and toggle views instantly from the navbar."),
    ("End-to-End Property Monetization Pipeline", "Sellers intake listings across districts, upload cadastral surveys and deeds, and select monetization packages (Basic, Verified, Featured) with PayHere integration."),
    ("LegalTech Deed Verification Workflow", "Title deeds uploaded by sellers are automatically routed to licensed real estate attorneys. Once cleared, listings earn the 'KNW Certified Deed Verified' badge to eliminate fraud."),
    ("Administrative Governance & Trust Gates", "Centralized moderation panel ensures no listing is published and no seller is approved without rigorous identity and legal document auditing.")
]

info_card_h = Inches(1.15)
for idx, (title, desc) in enumerate(arch_points):
    iy = Inches(1.5) + idx * (info_card_h + Inches(0.16))
    add_card(s12, info_x, iy, info_w, info_card_h, COLOR_CARD_BG, COLOR_BORDER)
    
    tb = s12.shapes.add_textbox(info_x + Inches(0.2), iy + Inches(0.12), info_w - Inches(0.4), info_card_h - Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    pt = tf.paragraphs[0]
    pt.text = f"{idx+1}. {title}"
    pt.font.name = FONT_HEADING
    pt.font.size = Pt(12)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_PRIMARY
    
    pd = tf.add_paragraph()
    pd.text = desc
    pd.font.name = FONT_BODY
    pd.font.size = Pt(10)
    pd.font.color.rgb = COLOR_BODY

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 13: BUYER & SELLER USE CASE WORKFLOWS
# ─────────────────────────────────────────────────────────────────────────────
s13 = prs.slides.add_slide(blank_layout)
add_header_and_footer(s13, "USE CASE SPECIFICATION", "Stakeholder Use Cases: Property Buyer & Seller Subsystems", 13)

sub_card_w = Inches(6.02)
sub_card_h = Inches(5.1)

# Left: Buyer
add_card(s13, Inches(0.5), Inches(1.5), sub_card_w, sub_card_h, COLOR_CARD_BG, COLOR_BORDER)
tb_b = s13.shapes.add_textbox(Inches(0.7), Inches(1.65), sub_card_w - Inches(0.4), Inches(0.4))
tf_b = tb_b.text_frame
tf_b.paragraphs[0].text = "👤 Property Buyer Use Case Model"
tf_b.paragraphs[0].font.name = FONT_HEADING
tf_b.paragraphs[0].font.size = Pt(14)
tf_b.paragraphs[0].font.bold = True
tf_b.paragraphs[0].font.color.rgb = COLOR_PRIMARY

if os.path.exists("diagrams/uc_buyer.png"):
    s13.shapes.add_picture("diagrams/uc_buyer.png", Inches(1.2), Inches(2.2), width=Inches(4.6))

# Right: Seller
add_card(s13, Inches(6.8), Inches(1.5), sub_card_w, sub_card_h, COLOR_CARD_BG, COLOR_BORDER)
tb_s = s13.shapes.add_textbox(Inches(7.0), Inches(1.65), sub_card_w - Inches(0.4), Inches(0.4))
tf_s = tb_s.text_frame
tf_s.paragraphs[0].text = "👤 Property Seller Use Case Model"
tf_s.paragraphs[0].font.name = FONT_HEADING
tf_s.paragraphs[0].font.size = Pt(14)
tf_s.paragraphs[0].font.bold = True
tf_s.paragraphs[0].font.color.rgb = COLOR_EMERALD

if os.path.exists("diagrams/uc_seller.png"):
    s13.shapes.add_picture("diagrams/uc_seller.png", Inches(7.5), Inches(2.3), width=Inches(4.6))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 14: LAWYER & ADMIN USE CASE WORKFLOWS
# ─────────────────────────────────────────────────────────────────────────────
s14 = prs.slides.add_slide(blank_layout)
add_header_and_footer(s14, "USE CASE SPECIFICATION", "Stakeholder Use Cases: Real Estate Lawyer & Administrator", 14)

# Left: Lawyer
add_card(s14, Inches(0.5), Inches(1.5), sub_card_w, sub_card_h, COLOR_CARD_BG, COLOR_BORDER)
tb_l = s14.shapes.add_textbox(Inches(0.7), Inches(1.65), sub_card_w - Inches(0.4), Inches(0.4))
tf_l = tb_l.text_frame
tf_l.paragraphs[0].text = "⚖️ Real Estate Lawyer Use Case Model"
tf_l.paragraphs[0].font.name = FONT_HEADING
tf_l.paragraphs[0].font.size = Pt(14)
tf_l.paragraphs[0].font.bold = True
tf_l.paragraphs[0].font.color.rgb = COLOR_AMBER

if os.path.exists("diagrams/uc_lawyer.png"):
    s14.shapes.add_picture("diagrams/uc_lawyer.png", Inches(1.2), Inches(2.3), width=Inches(4.6))

# Right: Admin
add_card(s14, Inches(6.8), Inches(1.5), sub_card_w, sub_card_h, COLOR_CARD_BG, COLOR_BORDER)
tb_a = s14.shapes.add_textbox(Inches(7.0), Inches(1.65), sub_card_w - Inches(0.4), Inches(0.4))
tf_a = tb_a.text_frame
tf_a.paragraphs[0].text = "🛡️ System Administrator Use Case Model"
tf_a.paragraphs[0].font.name = FONT_HEADING
tf_a.paragraphs[0].font.size = Pt(14)
tf_a.paragraphs[0].font.bold = True
tf_a.paragraphs[0].font.color.rgb = COLOR_TAG

if os.path.exists("diagrams/uc_admin.png"):
    s14.shapes.add_picture("diagrams/uc_admin.png", Inches(7.5), Inches(2.3), width=Inches(4.6))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 15: ENTITY-RELATIONSHIP (ER) DIAGRAM
# ─────────────────────────────────────────────────────────────────────────────
s15 = prs.slides.add_slide(blank_layout)
add_header_and_footer(s15, "DATABASE ARCHITECTURE", "Relational Database Schema (Entity-Relationship Diagram)", 15)

erd_w = Inches(12.33)
erd_h = Inches(5.1)
add_card(s15, Inches(0.5), Inches(1.5), erd_w, erd_h, COLOR_CARD_BG, COLOR_BORDER)

if os.path.exists("diagrams/er_diagram.png"):
    s15.shapes.add_picture("diagrams/er_diagram.png", Inches(1.5), Inches(1.65), width=Inches(10.33))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 16: UML CLASS DIAGRAM & DOMAIN ARCHITECTURE
# ─────────────────────────────────────────────────────────────────────────────
s16 = prs.slides.add_slide(blank_layout)
add_header_and_footer(s16, "OBJECT-ORIENTED DESIGN", "Domain Models, Business Logic & Controller Architecture", 16)

cd_w = Inches(12.33)
cd_h = Inches(5.1)
add_card(s16, Inches(0.5), Inches(1.5), cd_w, cd_h, COLOR_CARD_BG, COLOR_BORDER)

if os.path.exists("diagrams/class_diagram.png"):
    s16.shapes.add_picture("diagrams/class_diagram.png", Inches(2.2), Inches(1.65), width=Inches(9.0))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 17: SYSTEM SUMMARY & VIVA KEY TAKEAWAYS
# ─────────────────────────────────────────────────────────────────────────────
s17 = prs.slides.add_slide(blank_layout)
add_header_and_footer(s17, "CONCLUSION & EVALUATION", "Key Engineering Achievements & Viva Defense Takeaways", 17)

takeaways = [
    ("Unified Single-Account Engine", "Solved cross-role registration friction. Users operate seamlessly across Buyer, Seller, and Lawyer modes under one email address, with reactive Context API state synchronization."),
    ("Attorney-Backed LegalTech Security", "Eradicated fake listings by integrating real-estate attorneys directly into the property workflow. Uploaded deeds undergo rigorous legal audit prior to public marketplace certification."),
    ("Enterprise REST API Architecture", "Engineered with PHP 8.4 and Laravel 13, incorporating Bcrypt security, Eloquent ORM relations, private storage isolation, and modular controllers."),
    ("Full Commercial Integration", "Fully equipped with PayHere payment processing, dynamic mortgage/stamp duty calculators, automated ICS calendar sync, and interactive analytics dashboards.")
]

t_w = Inches(5.95)
t_h = Inches(2.35)

coords = [
    (Inches(0.5), Inches(1.5)),
    (Inches(6.88), Inches(1.5)),
    (Inches(0.5), Inches(4.2)),
    (Inches(6.88), Inches(4.2))
]

accents = [COLOR_PRIMARY, COLOR_EMERALD, COLOR_AMBER, COLOR_TAG]

for idx, (title, desc) in enumerate(takeaways):
    tx, ty = coords[idx]
    add_card(s17, tx, ty, t_w, t_h, COLOR_CARD_BG, COLOR_BORDER)
    
    # Left accent bar
    bar = s17.shapes.add_shape(MSO_SHAPE.RECTANGLE, tx, ty, Inches(0.1), t_h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accents[idx]
    bar.line.fill.background()
    
    tb = s17.shapes.add_textbox(tx + Inches(0.3), ty + Inches(0.2), t_w - Inches(0.5), t_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p1 = tf.paragraphs[0]
    p1.text = f"{idx+1}. {title}"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TITLE
    p1.space_after = Pt(8)
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.name = FONT_BODY
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = COLOR_BODY

# Save presentation
prs.save("KNW-PRO-Realtor.pptx")
print("Successfully added slides 11-17 to KNW-PRO-Realtor.pptx!")
